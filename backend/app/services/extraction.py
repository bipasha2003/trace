"""
TRACE - Text extraction from project source files

Runs at upload so the AI pipeline never re-parses binaries (PROJECT.md §11.2).
Supported: pdf, docx, md, txt. Anything else is refused at upload.
"""

import io
from pathlib import PurePath

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}
MAX_SOURCE_FILE_BYTES = 20 * 1024 * 1024  # 20 MB per file


class UnsupportedFile(ValueError):
    pass


def extension_of(file_name: str) -> str:
    return PurePath(file_name).suffix.lower()


def extract_text(file_name: str, data: bytes) -> str:
    ext = extension_of(file_name)
    if ext in {".md", ".markdown", ".txt"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
        except Exception as exc:  # pypdf raises a zoo of exception types for bad files
            raise UnsupportedFile(f"{file_name}: not a readable PDF ({exc.__class__.__name__})") from exc
    if ext == ".docx":
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise UnsupportedFile(f"{file_name}: not a readable .docx ({exc.__class__.__name__})") from exc
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    raise UnsupportedFile(
        f"{file_name}: unsupported type — upload one of {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )
